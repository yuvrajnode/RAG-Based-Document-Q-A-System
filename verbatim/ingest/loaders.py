"""File loaders.

Each loader turns a file into a list of Segments. Keeping page numbers on the
segment (rather than flattening to one blob) is what lets the UI say
"page 7 of the handbook" instead of just naming the file.
"""

from __future__ import annotations

import csv
import io
import json
import re
from pathlib import Path

from verbatim.logging import get_logger
from verbatim.models import LoadedDocument, Segment

log = get_logger(__name__)

SUPPORTED_EXTENSIONS = {
    ".pdf",
    ".docx",
    ".txt",
    ".md",
    ".markdown",
    ".html",
    ".htm",
    ".csv",
    ".json",
}

_WS = re.compile(r"[ \t ]+")
_BLANKS = re.compile(r"\n{3,}")
# Page headers/footers that repeat on every page add nothing but noise.
_PAGE_NUMBER_LINE = re.compile(r"^\s*(page\s+)?\d+\s*(/\s*\d+)?\s*$", re.IGNORECASE)


class UnsupportedFileType(ValueError):
    pass


def normalise(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = text.replace("­", "")  # soft hyphen
    # PDFs love to hyphenate across line breaks.
    text = re.sub(r"(\w)-\n(\w)", r"\1\2", text)
    lines = []
    for line in text.split("\n"):
        line = _WS.sub(" ", line).strip()
        if _PAGE_NUMBER_LINE.match(line):
            continue
        lines.append(line)
    return _BLANKS.sub("\n\n", "\n".join(lines)).strip()


def _load_pdf(path: Path) -> list[Segment]:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    segments: list[Segment] = []
    for number, page in enumerate(reader.pages, start=1):
        try:
            raw = page.extract_text() or ""
        except Exception as exc:  # a single broken page shouldn't kill the file
            log.warning("page %d of %s failed to extract: %s", number, path.name, exc)
            continue
        text = normalise(raw)
        if text:
            segments.append(Segment(text=text, page=number))
    return segments


def _load_docx(path: Path) -> list[Segment]:
    import docx

    document = docx.Document(str(path))
    segments: list[Segment] = []
    heading: str | None = None
    buffer: list[str] = []

    def flush() -> None:
        if buffer:
            segments.append(Segment(text=normalise("\n".join(buffer)), heading=heading))
            buffer.clear()

    for para in document.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        if para.style is not None and para.style.name.startswith("Heading"):
            flush()
            heading = text
            continue
        buffer.append(text)
    flush()

    for table in document.tables:
        rows = [" | ".join(cell.text.strip() for cell in row.cells) for row in table.rows]
        rendered = "\n".join(r for r in rows if r.strip(" |"))
        if rendered:
            segments.append(Segment(text=rendered, heading=heading))
    return [s for s in segments if s.text]


def _load_html(path: Path) -> list[Segment]:
    from bs4 import BeautifulSoup

    soup = BeautifulSoup(path.read_text(encoding="utf-8", errors="ignore"), "html.parser")
    for tag in soup(["script", "style", "nav", "footer", "noscript"]):
        tag.decompose()

    segments: list[Segment] = []
    heading: str | None = None
    for element in soup.find_all(["h1", "h2", "h3", "h4", "p", "li", "pre", "td"]):
        text = element.get_text(" ", strip=True)
        if not text:
            continue
        if element.name.startswith("h"):
            heading = text
            continue
        segments.append(Segment(text=normalise(text), heading=heading))
    if not segments:
        segments = [Segment(text=normalise(soup.get_text("\n")))]
    return [s for s in segments if s.text]


def _load_markdown(path: Path) -> list[Segment]:
    text = path.read_text(encoding="utf-8", errors="ignore")
    segments: list[Segment] = []
    heading: str | None = None
    buffer: list[str] = []
    in_code_block = False

    def flush() -> None:
        body = normalise("\n".join(buffer))
        if body:
            segments.append(Segment(text=body, heading=heading))
        buffer.clear()

    for line in text.split("\n"):
        if line.lstrip().startswith("```"):
            in_code_block = not in_code_block
            buffer.append(line)
            continue
        if not in_code_block and line.startswith("#"):
            flush()
            heading = line.lstrip("#").strip()
            continue
        buffer.append(line)
    flush()
    return segments


def _load_text(path: Path) -> list[Segment]:
    text = normalise(path.read_text(encoding="utf-8", errors="ignore"))
    return [Segment(text=text)] if text else []


def _load_csv(path: Path) -> list[Segment]:
    """Render each row as `column: value` lines - far more retrievable than raw CSV."""
    raw = path.read_text(encoding="utf-8", errors="ignore")
    reader = csv.DictReader(io.StringIO(raw))
    segments: list[Segment] = []
    for row_number, row in enumerate(reader, start=1):
        rendered = "\n".join(
            f"{key}: {value}" for key, value in row.items() if key and str(value).strip()
        )
        if rendered:
            segments.append(Segment(text=rendered, heading=f"Row {row_number}"))
    return segments


def _load_json(path: Path) -> list[Segment]:
    data = json.loads(path.read_text(encoding="utf-8", errors="ignore"))
    records = data if isinstance(data, list) else [data]
    segments = []
    for i, record in enumerate(records, start=1):
        segments.append(
            Segment(text=json.dumps(record, indent=2, ensure_ascii=False), heading=f"Record {i}")
        )
    return segments


_LOADERS = {
    ".pdf": _load_pdf,
    ".docx": _load_docx,
    ".html": _load_html,
    ".htm": _load_html,
    ".md": _load_markdown,
    ".markdown": _load_markdown,
    ".txt": _load_text,
    ".csv": _load_csv,
    ".json": _load_json,
}


def load_document(path: str | Path, *, title: str | None = None) -> LoadedDocument:
    path = Path(path)
    suffix = path.suffix.lower()
    loader = _LOADERS.get(suffix)
    if loader is None:
        raise UnsupportedFileType(
            f"{suffix or path.name!r} is not supported. "
            f"Supported: {', '.join(sorted(SUPPORTED_EXTENSIONS))}"
        )
    segments = loader(path)
    if not segments:
        raise ValueError(
            f"No extractable text in {path.name}. Scanned PDFs need OCR before ingestion."
        )
    return LoadedDocument(
        title=title or path.stem.replace("_", " ").replace("-", " ").strip(),
        source_type=suffix.lstrip("."),
        segments=segments,
        metadata={"filename": path.name, "bytes": path.stat().st_size},
    )
